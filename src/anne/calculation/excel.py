"""Excel workbook inspection primitives.

The extractor treats workbook formulas as candidate methodology evidence. It does
not declare the author's method correct; validation/backtesting is a later stage.
"""

from __future__ import annotations

from dataclasses import dataclass, field
import re
from pathlib import Path


@dataclass(frozen=True)
class CellFormula:
    sheet: str
    address: str
    formula: str
    dependencies: tuple[str, ...]


@dataclass
class WorkbookMethod:
    source: str
    formulas: list[CellFormula] = field(default_factory=list)
    inputs: list[str] = field(default_factory=list)
    outputs: list[str] = field(default_factory=list)


_CELL_RE = re.compile(r"(?<![A-Z0-9_])(?:'([^']+)'!)?\$?([A-Z]{1,3})\$?(\d+)")


def extract_formula_dependencies(formula: str) -> tuple[str, ...]:
    """Extract Excel-style cell references conservatively."""
    refs: list[str] = []
    for sheet, col, row in _CELL_RE.findall(formula.upper()):
        refs.append(f"{sheet}!{col}{row}" if sheet else f"{col}{row}")
    return tuple(dict.fromkeys(refs))


def inspect_workbook(path: str | Path) -> WorkbookMethod:
    """Inspect formulas using openpyxl when installed; otherwise fail explicitly."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel inspection requires the optional openpyxl dependency") from exc

    workbook = load_workbook(path, data_only=False, read_only=False)
    method = WorkbookMethod(source=str(Path(path)))
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    method.formulas.append(
                        CellFormula(
                            sheet=sheet.title,
                            address=cell.coordinate,
                            formula=cell.value,
                            dependencies=extract_formula_dependencies(cell.value),
                        )
                    )
    return method
